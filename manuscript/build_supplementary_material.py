#!/usr/bin/env python3
"""Build the submission-ready Supplementary Tables S1--S7 package."""

from __future__ import annotations

import os
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt


HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "submission_files"
TABLES = tuple(
    HERE / f"supplementary_table_S{number}_{name}.md"
    for number, name in (
        (1, "arcgis_planning_objectives"),
        (2, "neighbourhood_weight_sensitivity"),
        (3, "flus_feature_diagnostics"),
        (4, "cross_product_class_matrices"),
        (5, "rolling_external_diagnostics"),
        (6, "product_robustness"),
        (7, "allocation_limits"),
    )
)
MARKDOWN_FORMAT = "markdown+pipe_tables+tex_math_dollars"


def run(*args: str) -> None:
    subprocess.run(args, cwd=HERE, check=True)


def source_text() -> str:
    sections = []
    for table in TABLES:
        if not table.is_file():
            raise FileNotFoundError(table)
        sections.append(table.read_text(encoding="utf-8").strip())
    return "\n\n\\newpage\n\n".join(sections) + "\n"


def style_docx(path: Path) -> None:
    document = Document(path)
    for section in document.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297)
        section.page_height = Mm(210)
        section.left_margin = section.right_margin = Mm(15)
        section.top_margin = section.bottom_margin = Mm(15)
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(8)
    for table in document.tables:
        table.autofit = True
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(6.5)
    document.save(path)


def build() -> None:
    OUTPUT.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="lup-supplement-") as directory:
        source = Path(directory) / "Supplementary_Material.md"
        source.write_text(source_text(), encoding="utf-8")
        pdf = OUTPUT / "Supplementary_Material.pdf"
        docx = OUTPUT / "Supplementary_Material.docx"
        run(
            "pandoc",
            str(source),
            f"--from={MARKDOWN_FORMAT}",
            "--standalone",
            f"--pdf-engine={os.environ.get('MANUSCRIPT_TEX_ENGINE', 'xelatex')}",
            f"--variable=mainfont:{os.environ.get('MANUSCRIPT_FONT', 'Arial Unicode MS')}",
            "--variable=papersize:a4",
            "--variable=classoption:landscape",
            "--variable=fontsize:8pt",
            "--variable=geometry:top=15mm,bottom=15mm,left=15mm,right=15mm",
            "--output",
            str(pdf),
        )
        run(
            "pandoc",
            str(source),
            f"--from={MARKDOWN_FORMAT}",
            "--standalone",
            "--output",
            str(docx),
        )
        style_docx(docx)


if __name__ == "__main__":
    build()
