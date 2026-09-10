#!/usr/bin/env python3
"""Build synchronized author and blinded manuscript submission artifacts."""

from __future__ import annotations

import shutil
import subprocess
import os
from pathlib import Path


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SUBMISSION = HERE / "submission_files"
MARKDOWN_FORMAT = (
    "markdown+yaml_metadata_block+implicit_figures+tex_math_dollars+"
    "pipe_tables+link_attributes"
)
TITLE = "Auditing constrained geospatial allocation under annual land-cover product uncertainty"


def compile_tex(source: str, cwd: Path = HERE) -> None:
    engine = os.environ.get("MANUSCRIPT_TEX_ENGINE", "xelatex")
    if engine == "tectonic":
        run(engine, "--keep-logs", source, cwd=cwd)
    else:
        for _ in range(2):
            run(engine, "-interaction=nonstopmode", "-halt-on-error", source, cwd=cwd)


def run(*args: str, cwd: Path = HERE) -> None:
    subprocess.run(args, cwd=cwd, check=True)


def markdown_to_latex(source: Path, target: Path) -> None:
    run(
        "pandoc",
        str(source),
        f"--from={MARKDOWN_FORMAT}",
        "--to=latex",
        f"--resource-path={ROOT}",
        "--output",
        str(target),
    )
    content = target.read_text(encoding="utf-8")
    if content.startswith("\\maketitle\n"):
        content = content.removeprefix("\\maketitle\n")
    target.write_text(content.rstrip() + "\n", encoding="utf-8")


def markdown_to_docx(source: Path, target: Path) -> None:
    raw_target = target.with_name(f".{target.stem}.pandoc.docx")
    ready_target = target.with_name(f".{target.stem}.ready.docx")
    run(
        "pandoc",
        str(source),
        f"--from={MARKDOWN_FORMAT}",
        "--standalone",
        f"--resource-path={ROOT}",
        "--output",
        str(raw_target),
    )
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    document = Document(raw_target)
    if source.resolve() == (HERE / "manuscript.md").resolve():
        first = document.paragraphs[0]
        first.insert_paragraph_before(TITLE, style="Title")
        first.insert_paragraph_before("Ning Zhou\nBeijing Freedo Technology Co., Ltd.\nCorresponding author: zhouning@freedotech.com", style="Subtitle")
    for section in document.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(25)
        section.top_margin = section.bottom_margin = Mm(25)
    for name in ("Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        if name in document.styles:
            document.styles[name].font.name = "Times New Roman"
            document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.styles["Normal"].font.size = Pt(11)
    figure_index = 0
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Captioned Figure":
            paragraph.paragraph_format.keep_with_next = True
        elif paragraph.style.name == "Image Caption":
            paragraph.paragraph_format.keep_together = True
            figure_index += 1
            if paragraph.runs:
                paragraph.runs[0].text = f"Figure {figure_index}. " + paragraph.runs[0].text
            else:
                paragraph.add_run(f"Figure {figure_index}. ")
    for shape in document.inline_shapes:
        if shape.width > Mm(160):
            scale = Mm(160) / shape.width
            shape.width = Mm(160)
            shape.height = int(shape.height * scale)
    document.save(ready_target)
    os.replace(ready_target, target)
    raw_target.unlink()


def markdown_to_pdf(source: Path, target: Path) -> None:
    run(
        "pandoc",
        str(source),
        f"--from={MARKDOWN_FORMAT}",
        "--standalone",
        f"--pdf-engine={os.environ.get('MANUSCRIPT_TEX_ENGINE', 'xelatex')}",
        f"--resource-path={ROOT}",
        f"--variable=mainfont:{os.environ.get('MANUSCRIPT_FONT', 'Arial Unicode MS')}",
        f"--include-in-header={HERE / 'submission_header.tex'}",
        "--variable=papersize:a4",
        "--variable=fontsize:10pt",
        "--variable=geometry:top=22mm,bottom=20mm,left=20mm,right=20mm",
        "--output",
        str(target),
    )


def anonymize(main: str) -> str:
    body, references = main.split("## Data availability", maxsplit=1)
    references = "## References" + references.split("## References", maxsplit=1)[1]
    body = body.replace(
        "the accompanying source repository `FLUS_console_crossplatform` at commit `47e65b3` contains the author modifications used for this paper.",
        "an archived source repository at an anonymized revision contains the modifications used for this paper.",
    )
    body = body.replace("the author's `train`/`train-update`", "the modified `train`/`train-update`")
    return "\n".join(
        [
            "---",
            f'title: "{TITLE}"',
            'author: ""',
            'date: ""',
            "---",
            "",
            body.strip(),
            "",
            "## Data availability",
            "",
            "The public-data benchmark manifests, reports, generated delivery artefacts and source-data tables are available from an anonymized repository and archival package. Repository and archive identifiers are withheld during double-anonymized review and will be supplied in the accepted version. The package includes the Dynamic World and ArcGIS-served product-track reports, aligned public rasters, planning outputs, raster-cell transition footprints and integrity manifests. Authoritative Abu Dhabi land-use and independent reference data are not part of this study.",
            "",
            "## Code availability",
            "",
            "The analysis scripts, benchmark protocol, model runners, figure-generation code, model assets and integrity checks are maintained in the accompanying anonymized repository and archival package. Identifiers and contact details are withheld during double-anonymized review and will be supplied in the accepted version.",
            "",
            "## Declarations",
            "",
            "Funding, competing interests, author contributions, acknowledgements and ethics information are supplied in the separate title page and submission metadata.",
            "",
            references.strip(),
            "",
        ]
    )


def submission_texts(main: str) -> dict[str, str]:
    abstract = main.split("## Abstract", maxsplit=1)[1].split("## Highlights", maxsplit=1)[0].strip()
    highlights = main.split("## Highlights", maxsplit=1)[1].split("## Keywords", maxsplit=1)[0].strip()
    cover = f"""# Cover Letter

Dear Editor,

Please consider our manuscript entitled \"{TITLE}\" for publication as a Research Article in *Landscape and Urban Planning*.

The manuscript evaluates Geospatial Kernel, the algorithmic core of a Geospatial World Model, as an auditable execution layer for constrained spatial allocation. Using two public annual land-cover product tracks for Abu Dhabi, it tests whether model rankings are robust to product choice while retaining explicit proposal, constraint-projection and state-writeback traces.

The study is deliberately framed as a public-data benchmark and conditional planning stress test, rather than as an official forecast of Abu Dhabi statutory land-use change. The refreshed source release, including the ArcGIS-served product track, is archived at https://doi.org/10.5281/zenodo.22689057; hydrated Git LFS artifacts are retrieved from the corresponding `v2.0.0` repository tag.

I confirm that the manuscript is original, is not under consideration by another journal, and has been approved for submission.

The corresponding author is:

Ning Zhou\\
Beijing Freedo Technology Co., Ltd.\\
Email: zhouning@freedotech.com

Sincerely,

Ning Zhou
"""
    title_page = f"""# Title Page

## Title

{TITLE}

## Author

Ning Zhou

## Affiliation

Beijing Freedo Technology Co., Ltd.

## Corresponding author

Ning Zhou\\
Beijing Freedo Technology Co., Ltd.\\
Email: zhouning@freedotech.com

## Declarations

**Funding.** No external funding was received for this study.

**Competing interests.** Ning Zhou is employed by Beijing Freedo Technology Co., Ltd., which develops geospatial-world-model software. The author declares no other competing financial or personal interests.

**CRediT author statement.** Ning Zhou: Conceptualization, methodology, software, data curation, formal analysis, visualization, writing—original draft, writing—review and editing.

**Acknowledgements.** The author thanks the anonymous reviewers for independent technical verification and constructive methodological comments.

**Ethics statement.** Not applicable; the study used geospatial raster, vector and derived public-data products and did not involve human or animal participants.

**Data and code.** The ArcGIS-served public benchmark source release is archived at https://doi.org/10.5281/zenodo.22689057 (v2.0.0; GitHub tag `v2.0.0`, commit `77e2c59`; concept DOI https://doi.org/10.5281/zenodo.22663475). The Zenodo-generated ZIP preserves Git LFS pointer identities; hydrated large artifacts required for a full numerical rerun are retrieved from the tagged repository and verified against the included manifests.
"""
    return {
        "abstract.md": f"# Abstract\n\n{abstract}\n",
        "highlights.md": f"# Highlights\n\n{highlights}\n",
        "cover_letter.md": cover,
        "title_page.md": title_page,
    }


def build() -> None:
    main_markdown = HERE / "manuscript.md"
    main = main_markdown.read_text(encoding="utf-8")
    SUBMISSION.mkdir(exist_ok=True)

    markdown_to_latex(main_markdown, HERE / "manuscript_pandoc.tex")
    compile_tex("main.tex")
    compile_tex("lup_submission.tex")
    shutil.copy2(HERE / "main.pdf", HERE / "manuscript.pdf")
    markdown_to_docx(main_markdown, HERE / "manuscript.docx")

    anonymous_markdown = SUBMISSION / "manuscript_anonymous.md"
    anonymous_markdown.write_text(anonymize(main), encoding="utf-8")
    markdown_to_latex(anonymous_markdown, SUBMISSION / "manuscript_anonymous_pandoc.tex")
    compile_tex("manuscript_anonymous.tex", cwd=SUBMISSION)
    markdown_to_docx(anonymous_markdown, SUBMISSION / "manuscript_anonymous.docx")

    for markdown_name, content in submission_texts(main).items():
        source = SUBMISSION / markdown_name
        source.write_text(content, encoding="utf-8")
        markdown_to_docx(source, source.with_suffix(".docx"))
        markdown_to_pdf(source, source.with_suffix(".pdf"))


if __name__ == "__main__":
    build()
